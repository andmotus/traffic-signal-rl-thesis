"""
Section 4.5 and H4: measured CO2 by vehicle class and the national scaling model.

Reproduces
    Table 4.4   CO2 per vehicle by vehicle class, Fixed-Time compared with standard-state IDQN
    Section 4.5 representativeness of the two CO2 validation runs
    Figure 4.4  CO2 per vehicle by vehicle class
    Table 3.4   inputs of the national scaling model (Steps 1-4 of Section 3.7)
    Table 4.5   national CO2-reduction potential, with the sensitivities of Section 4.5
    Figure 4.5  national CO2-reduction potential by scenario and city type

CO2 comes from the SUMO tripinfo file of the final episode of each validation
run (Fixed-Time: episode 100, IDQN: episode 200). Vehicle classes follow the
KBA definition used for the injection: the InTAS buses count as trucks.

Inputs that are not files in this repository are constants below, each with
its source:
    - published OpenStreetMap node counts (Section 3.7)
    - fixed-time pool, adoption and peak-hour equivalents (Table 3.4)
    - German transport-sector emissions (Umweltbundesamt, 2026)
    - the plain-SUMO baseline measurement used as a plausibility check
"""
import csv
import re
import statistics
import xml.etree.ElementTree as ET

import numpy as np
import openpyxl

from common import (BLUE, DOCS, FIGURES, GRAY, NETWORKS, ORANGE, ROOT, converged_values, emissions_run,
                    fmt, markdown_table, plot_style, print_utf8, relative_change, save_csv, series,
                    training_runs)

CLASSES = ["car", "truck", "motorcycle"]

# --- Step 3: signalised intersections -------------------------------------------------------
# Nodes tagged highway=traffic_signals within the city boundary (OpenStreetMap contributors, 2026),
# as reported in Section 3.7. Entered as constants: the Overpass query itself is not yet in the repository.
OSM_SIGNAL_NODES = {"cologne": 3568, "ingolstadt": 558}

# --- Step 4: baseline and adoption (Table 3.4) ----------------------------------------------
FIXED_TIME_POOL = {"cologne": (0.14, 0.205, 0.27),      # Fakler et al. (2014); Bodenheimer et al. (2014)
                   "ingolstadt": (0.11, 0.175, 0.24)}
ADOPTION = (0.25, 0.50, 1.00)                            # scenario assumption
PEAK_HOUR_EQUIVALENTS = (2, 4, 8)                        # per day, scenario assumption
DAYS_PER_YEAR = 365
SCENARIOS = ("Low", "Medium", "High")

# --- Step 5: reference -------------------------------------------------------------------
# German transport sector 2025: 146.3 million t CO2-equivalent (Umweltbundesamt, 2026)
TRANSPORT_EMISSIONS_T = 146.3e6
TRANSPORT_EMISSIONS_LABEL = "146.3 million t CO2e, transport sector 2025 (Umweltbundesamt, 2026)"

# Fixed-Time on Cologne1 measured with plain SUMO, outside RESCO, at the start of the project (Section 4.5)
PLAIN_SUMO_BASELINE_G = 175.7

KBA_FILE = DOCS / "fz1_2024.xlsx"
DESTATIS_FILE = DOCS / "06-staedte-einwohner-groessen.xlsx"
COLOGNE_INVENTORY = DOCS / "lsaknoten_20241031.csv"
KOELN, INGOLSTADT = "05315", "09161"
FREIBURG = "08311"


# ================================================================ CO2 per vehicle class

def vtype_classes(network):
    """vType id -> vClass from the typed route file of a network."""
    root = ET.parse(ROOT / "RESCO" / "resco_benchmark" / "environments" / network / f"{network}_typed.rou.xml").getroot()
    return {vt.get("id"): vt.get("vClass", "passenger") for vt in root.iter("vType")}


def vehicle_class(vtype, vclasses):
    if vtype == "truck" or vclasses.get(vtype) == "bus":
        return "truck"
    if vtype == "motorcycle":
        return "motorcycle"
    return "car"


def final_tripinfo(network, controller):
    folder = emissions_run(network, controller).with_suffix("")
    folder = folder.parent / folder.name[:-8]  # run folder = JSON name without the last 8 hash characters
    files = sorted(folder.rglob("tripinfo_*.xml"), key=lambda p: int(p.stem.split("_")[1]))
    if not files:
        raise FileNotFoundError(f"no tripinfo file under {folder}")
    return files[-1]


def co2_by_class(network, controller):
    """{class: (vehicles, mean CO2 per vehicle in g)} plus 'all'."""
    vclasses = vtype_classes(network)
    sums = {c: [0.0, 0] for c in CLASSES}
    for trip in ET.parse(final_tripinfo(network, controller)).getroot().iter("tripinfo"):
        em = trip.find("emissions")
        c = vehicle_class(trip.get("vType"), vclasses)
        sums[c][0] += float(em.get("CO2_abs")) / 1000.0  # mg -> g
        sums[c][1] += 1
    out = {c: (n, total / n) for c, (total, n) in sums.items()}
    n_all = sum(n for n, _ in out.values())
    out["all"] = (n_all, sum(total for total, _ in sums.values()) / n_all)
    return out


def table_4_4():
    rows, csv_rows, savings, vehicles = [], [], {}, {}
    for net, label in NETWORKS.items():
        ft, idqn = co2_by_class(net, "fixed"), co2_by_class(net, "idqn")
        savings[net] = {}
        for c in CLASSES + ["all"]:
            (n_ft, g_ft), (n_rl, g_rl) = ft[c], idqn[c]
            assert n_ft == n_rl, f"{net} {c}: different vehicle counts ({n_ft} vs {n_rl})"
            name = {"car": "Car", "truck": "Truck", "motorcycle": "Motorcycle", "all": "All vehicles"}[c]
            if c == "truck" and net == "ingolstadt1":
                name = "Truck (incl. buses)"
            rows.append([label if c == "car" else "", name, f"{n_ft:,}", fmt(g_ft, 1), fmt(g_rl, 1),
                         fmt(relative_change(g_rl, g_ft), 1, pct=True)])
            csv_rows.append([net, c, n_ft, g_ft, g_rl, g_ft - g_rl, relative_change(g_rl, g_ft)])
            if c != "all":
                savings[net][c] = g_ft - g_rl
        vehicles[net] = ft["all"][0]
    save_csv("table_4_4_co2_by_class", ["network", "class", "vehicles", "fixed_g", "idqn_g", "saving_g", "change"],
             csv_rows)
    header = ["Network", "Class", "Vehicles", "Fixed-Time (g)", "IDQN (g)", "Change"]
    return markdown_table(header, rows), savings, vehicles


def representativeness():
    lines = []
    runs = training_runs("multiseed21")
    for net, label in NETWORKS.items():
        final_rl = series(emissions_run(net, "idqn"), "delay")[-1]
        final_ft = series(emissions_run(net, "fixed"), "delay")[-1]
        conv = converged_values(runs[(net, "A")], "delay")
        lines.append(f"- {label}: final IDQN episode delay {fmt(final_rl, 2)} s "
                     f"(median converged delay of the 21 runs {fmt(np.median(conv), 2)} s; "
                     f"{int((conv < final_rl).sum())} of 21 runs better); final Fixed-Time episode "
                     f"{fmt(final_ft, 2)} s, so the IDQN episode is {fmt(relative_change(final_rl, final_ft), 1, pct=True)}")
    ft_all = co2_by_class("cologne1", "fixed")["all"][1]
    lines.append(f"- Plausibility: Cologne1 Fixed-Time {fmt(ft_all, 1)} g per vehicle against "
                 f"{PLAIN_SUMO_BASELINE_G} g with plain SUMO ({fmt(relative_change(ft_all, PLAIN_SUMO_BASELINE_G), 1, pct=True)})")
    return "\n".join(lines)


def figure_4_4(plt):
    fig, axes = plt.subplots(1, 2, figsize=(6.3, 2.5), sharey=False)
    fig.subplots_adjust(left=0.08, right=0.99, top=0.9, bottom=0.13, wspace=0.18)
    for ax, (net, label) in zip(axes, NETWORKS.items()):
        ft, rl = co2_by_class(net, "fixed"), co2_by_class(net, "idqn")
        x = np.arange(3)
        a = [ft[c][1] for c in CLASSES]
        b = [rl[c][1] for c in CLASSES]
        ax.bar(x - 0.18, a, 0.36, color=GRAY, label="Fixed-Time")
        ax.bar(x + 0.18, b, 0.36, color=BLUE, label="IDQN, standard")
        for k in range(3):
            ax.text(x[k] + 0.18, b[k] + 3, fmt(relative_change(b[k], a[k]), 1, pct=True).replace("−", "-"),
                    ha="center", va="bottom", fontsize=7.5)
        ax.set_xticks(x)
        ax.set_xticklabels(["Car", "Truck¹", "Motorcycle"])
        ax.set_ylim(0, 320)
        ax.set_title(label)
        if net == "cologne1":
            ax.set_ylabel("CO$_2$ per vehicle (g)")
            ax.legend(frameon=False, loc="upper left", fontsize=7.5)
    fig.savefig(FIGURES / "Figure_4_4_co2_by_class.png")
    plt.close(fig)


# ================================================================ national scaling model

def kba_districts():
    """Registration districts of KBA FZ 1.1 (1 January 2024) with derived population and class counts.

    Truck = Kraftomnibusse + Lastkraftwagen insgesamt + Zugmaschinen insgesamt (Table 3.1).
    Population = Pkw / (Pkw-Dichte je 1.000 Einwohner) * 1,000.
    """
    ws = openpyxl.load_workbook(KBA_FILE, read_only=True, data_only=True)["FZ1.1"]
    rows = list(ws.iter_rows(values_only=True))
    top, sub = rows[7], rows[8]
    group, cols = None, {}
    for i, (t, s) in enumerate(zip(top, sub)):
        group = (t or group or "").strip() if t else group
        cols[(re.sub(r"\s+", " ", group or "").strip(), re.sub(r"\s+", " ", s or "").strip())] = i
    c_moto = cols[("Krafträder", "insgesamt")]
    c_pkw = cols[("Personenkraftwagen", "insgesamt")]
    c_dens = cols[("Personenkraftwagen", "PKW-Dichte je 1.000 Einwohner")]
    c_bus = cols[("Kraftomni- busse", "")]
    c_lkw = cols[("Lastkraftwagen", "insgesamt")]
    c_zug = cols[("Zugmaschinen", "insgesamt")]
    num = lambda v: 0 if v in (None, "-", ".") else v
    out = []
    for r in rows[9:]:
        name = r[3]
        if not (isinstance(name, str) and re.match(r"^\d{5} ", name)):
            continue
        pkw, moto = num(r[c_pkw]), num(r[c_moto])
        truck = num(r[c_bus]) + num(r[c_lkw]) + num(r[c_zug])
        total = pkw + moto + truck
        out.append(dict(code=name[:5], name=name[6:].strip(), pkw=pkw, moto=moto, truck=truck,
                        population=pkw / num(r[c_dens]) * 1000,
                        share={"car": pkw / total, "truck": truck / total, "motorcycle": moto / total}))
    return out


def city_name(d):
    return re.sub(r",\s*STADT$", "", d["name"]).title()


def is_kreisfrei(d, include_freiburg=False):
    """Districts that are a single city (kreisfreie Stadt): KBA names ending in ',STADT' plus Berlin and Hamburg.

    This is the rule used for the thesis. It misses Freiburg im Breisgau, whose KBA name is
    'FREIBURG I.BREISG.STADT' (no comma); include_freiburg=True adds it as a sensitivity.
    """
    if d["code"] in ("11000", "02000"):
        return True
    if include_freiburg and d["code"] == FREIBURG:
        return True
    return bool(re.search(r",\s*STADT$", d["name"]))


def destatis_city_counts():
    """Number of cities with 100,000-500,000 and with 500,000+ inhabitants on 31 December 2024."""
    ws = openpyxl.load_workbook(DESTATIS_FILE, read_only=True, data_only=True)["Städte Einwohnergrößenklassen"]
    rows = list(ws.iter_rows(values_only=True))
    header = rows[1]
    c_mid = next(i for i, h in enumerate(header) if h and "100 000 bis unter 500 000" in h)
    c_big = next(i for i, h in enumerate(header) if h and "500 000 und mehr" in h)
    total = next(r for r in rows if r[1] and str(r[1]).strip() == "Deutschland")
    return {"ingolstadt": int(total[c_mid]), "cologne": int(total[c_big])}


def city_types(districts, include_freiburg=False):
    """Kreisfreie cities per city type, median composition and population of the type (Steps 2-3)."""
    counts = destatis_city_counts()
    kf = [d for d in districts if is_kreisfrei(d, include_freiburg)]
    groups = {"cologne": [d for d in kf if d["population"] >= 500_000],
              "ingolstadt": [d for d in kf if 100_000 <= d["population"] < 500_000]}
    out = {}
    for t, cities in groups.items():
        truck = statistics.median(d["share"]["truck"] for d in cities)
        moto = statistics.median(d["share"]["motorcycle"] for d in cities)
        pops = [d["population"] for d in cities]
        missing = max(0, counts[t] - len(cities))
        out[t] = dict(n_official=counts[t], n_kreisfrei=len(cities), n_imputed=missing,
                      composition={"car": 1 - truck - moto, "truck": truck, "motorcycle": moto},
                      population=sum(pops) + missing * statistics.median(pops))
    return out


def intersection_rates(districts):
    """Signalised intersections per 1,000 inhabitants of the two calibration cities (Step 3)."""
    with open(COLOGNE_INVENTORY, encoding="utf-8-sig") as f:
        cologne_official = sum(1 for _ in csv.DictReader(f))
    nodes_per_intersection = OSM_SIGNAL_NODES["cologne"] / cologne_official
    ingolstadt_est = round(OSM_SIGNAL_NODES["ingolstadt"] / nodes_per_intersection)
    pop = {d["code"]: d["population"] for d in districts}
    return dict(cologne_official=cologne_official, nodes_per_intersection=nodes_per_intersection,
                ingolstadt_estimated=ingolstadt_est,
                rate={"cologne": cologne_official / pop[KOELN] * 1000,
                      "ingolstadt": ingolstadt_est / pop[INGOLSTADT] * 1000})


def national_model(savings, vehicles, include_freiburg=False, composition=None):
    districts = kba_districts()
    types = city_types(districts, include_freiburg)
    rates = intersection_rates(districts)
    net_of = {"cologne": "cologne1", "ingolstadt": "ingolstadt1"}
    model = {}
    for t, info in types.items():
        comp = composition[t] if composition else info["composition"]
        net = net_of[t]
        saving_kg = vehicles[net] * sum(comp[c] * savings[net][c] for c in CLASSES) / 1000
        n_int = rates["rate"][t] * info["population"] / 1000
        annual = [n_int * pool * adopt * saving_kg * phe * DAYS_PER_YEAR / 1000
                  for pool, adopt, phe in zip(FIXED_TIME_POOL[t], ADOPTION, PEAK_HOUR_EQUIVALENTS)]
        all_signals = [n_int * adopt * saving_kg * phe * DAYS_PER_YEAR / 1000
                       for adopt, phe in zip(ADOPTION, PEAK_HOUR_EQUIVALENTS)]
        model[t] = dict(info, composition=comp, saving_kg=saving_kg, rate=rates["rate"][t],
                        intersections=n_int, annual_t=annual, all_signals_t=all_signals)
    return model, rates, districts


def table_3_4(model, rates, savings, vehicles):
    c, i = model["cologne"], model["ingolstadt"]
    pct = lambda comp: " / ".join(f"{100 * comp[k]:.2f}" for k in CLASSES) + "%"
    g = lambda net: " / ".join(f"{savings[net][k]:.2f}" for k in CLASSES)
    adoption = " / ".join(f"{100 * a:g}" for a in ADOPTION) + "%"
    phe = " / ".join(map(str, PEAK_HOUR_EQUIVALENTS))
    rows = [
        ["Vehicles per peak hour", f"{vehicles['cologne1']:,}", f"{vehicles['ingolstadt1']:,}"],
        ["CO2 saving per vehicle: car / truck / motorcycle (g)", g("cologne1"), g("ingolstadt1")],
        ["Fleet composition: car / truck / motorcycle (median of kreisfreie cities)", pct(c["composition"]), pct(i["composition"])],
        ["Saving per intersection and peak hour", f"{c['saving_kg']:.1f} kg", f"{i['saving_kg']:.1f} kg"],
        ["Number of cities (Destatis, 31 Dec 2024)", c["n_official"], i["n_official"]],
        ["Kreisfreie cities found in KBA FZ 1 / imputed with the median population", f"{c['n_kreisfrei']} / {c['n_imputed']}", f"{i['n_kreisfrei']} / {i['n_imputed']}"],
        ["Population of the city type (million)", f"{c['population'] / 1e6:.2f}", f"{i['population'] / 1e6:.2f}"],
        ["Signalised intersections per 1,000 inhabitants", f"{c['rate']:.4f}", f"{i['rate']:.4f}"],
        ["Signalised intersections (estimated)", f"{c['intersections']:,.0f}", f"{i['intersections']:,.0f}"],
        ["Fixed-time pool (low / central / high)", " / ".join(f"{100 * p:g}" for p in FIXED_TIME_POOL["cologne"]) + "%",
         " / ".join(f"{100 * p:g}" for p in FIXED_TIME_POOL["ingolstadt"]) + "%"],
        ["Adoption within pool", adoption, adoption],
        ["Peak-hour equivalents per day", phe, phe],
    ]
    save_csv("table_3_4_inputs", ["input", "cologne_type", "ingolstadt_type"], rows)
    notes = (f"Cologne: {rates['cologne_official']} intersections in the Stadt Koeln inventory; OpenStreetMap "
             f"{OSM_SIGNAL_NODES['cologne']:,} nodes, {rates['nodes_per_intersection']:.2f} per intersection; Ingolstadt "
             f"{OSM_SIGNAL_NODES['ingolstadt']} nodes -> {rates['ingolstadt_estimated']} intersections.")
    return markdown_table(["Input", "Cologne type (> 500,000)", "Ingolstadt type (100,000–500,000)"], rows) + "\n\n" + notes


def table_4_5(model):
    c, i = model["cologne"], model["ingolstadt"]
    rows, csv_rows = [], []
    for k, name in enumerate(SCENARIOS):
        total = c["annual_t"][k] + i["annual_t"][k]
        rows.append([name, f"{c['annual_t'][k]:,.0f}", f"{i['annual_t'][k]:,.0f}", f"{total:,.0f}",
                     f"{100 * total / TRANSPORT_EMISSIONS_T:.3f}%"])
        csv_rows.append([name, c["annual_t"][k], i["annual_t"][k], total, total / TRANSPORT_EMISSIONS_T,
                         i["annual_t"][k] / total])
    sens = [c["all_signals_t"][k] + i["all_signals_t"][k] for k in range(3)]
    rows.append(["Sensitivity: all signals replaceable, low / medium / high", "", "",
                 " / ".join(f"{v:,.0f}" for v in sens), " / ".join(f"{100 * v / TRANSPORT_EMISSIONS_T:.3f}" for v in sens) + "%"])
    for k, name in enumerate(SCENARIOS):
        csv_rows.append([f"{name} (all signals)", c["all_signals_t"][k], i["all_signals_t"][k], sens[k],
                         sens[k] / TRANSPORT_EMISSIONS_T, i["all_signals_t"][k] / sens[k]])
    save_csv("table_4_5_national", ["scenario", "cologne_type_t", "ingolstadt_type_t", "total_t",
                                    "share_of_transport", "ingolstadt_type_share"], csv_rows)
    header = ["Scenario", "Cologne type (t/yr)", "Ingolstadt type (t/yr)", "Total (t/yr)", "Share of transport CO2"]
    return markdown_table(header, rows), sens


def sensitivities(model, savings, vehicles, districts):
    base = [model["cologne"]["annual_t"][k] + model["ingolstadt"]["annual_t"][k] for k in range(3)]
    shares = [model["ingolstadt"]["annual_t"][k] / base[k] for k in range(3)]
    lines = ["- Share of the Ingolstadt type in the total: " + " / ".join(f"{100 * s:.1f}%" for s in shares),
             f"- Signalised intersections in the 79 cities: {model['cologne']['intersections'] + model['ingolstadt']['intersections']:,.0f}",
             f"- Range high / low: factor {base[2] / base[0]:.1f}"]
    # composition of the simulated demand (vehicle counts of the typed Cologne1 and Ingolstadt1 demand)
    own = {}
    for t, net in [("cologne", "cologne1"), ("ingolstadt", "ingolstadt1")]:
        counts = co2_by_class(net, "fixed")
        own[t] = {c: counts[c][0] / counts["all"][0] for c in CLASSES}
    alt, _, _ = national_model(savings, vehicles, composition=own)
    alt_tot = [alt["cologne"]["annual_t"][k] + alt["ingolstadt"]["annual_t"][k] for k in range(3)]
    lines.append("- Composition of the simulated demand instead of the city-type medians: "
                 + " / ".join(fmt(relative_change(a, b), 2, sign=True, pct=True) for a, b in zip(alt_tot, base)))
    # Freiburg im Breisgau counted as a kreisfreie city
    fr, _, _ = national_model(savings, vehicles, include_freiburg=True)
    fr_tot = [fr["cologne"]["annual_t"][k] + fr["ingolstadt"]["annual_t"][k] for k in range(3)]
    lines.append(f"- Freiburg im Breisgau included as a kreisfreie city ({fr['ingolstadt']['n_kreisfrei']} found, "
                 f"{fr['ingolstadt']['n_imputed']} imputed; composition "
                 + " / ".join(f"{100 * fr['ingolstadt']['composition'][k]:.2f}" for k in CLASSES)
                 + "%): totals " + " / ".join(f"{v:,.0f}" for v in fr_tot) + " t ("
                 + " / ".join(fmt(relative_change(a, b), 2, sign=True, pct=True) for a, b in zip(fr_tot, base)) + ")")
    # registration-seat outlier mentioned in Section 3.7
    for t, (lo, hi) in [("cologne", (500_000, float("inf"))), ("ingolstadt", (100_000, 500_000))]:
        kf = sorted((d for d in districts if is_kreisfrei(d) and lo <= d["population"] < hi),
                    key=lambda d: -d["share"]["truck"])
        lines.append(f"- Highest truck shares, {t.title()} type: "
                     + ", ".join(f"{city_name(d)} {100 * d['share']['truck']:.1f}%" for d in kf[:3]))
    return "\n".join(lines)


def figure_4_5(model, plt):
    c = np.array(model["cologne"]["annual_t"]) / 1000
    i = np.array(model["ingolstadt"]["annual_t"]) / 1000
    tot = c + i
    fig, ax = plt.subplots(figsize=(6.3, 3.0))
    fig.subplots_adjust(left=0.11, right=0.98, top=0.95, bottom=0.12)
    x = np.arange(3)
    ax.bar(x, c, 0.5, color=BLUE, label="Cologne-type cities (> 500,000)", edgecolor="white", lw=1)
    ax.bar(x, i, 0.5, bottom=c, color=ORANGE, label="Ingolstadt-type cities (100,000–500,000)", edgecolor="white", lw=1)
    for k in range(3):
        ax.text(x[k], tot[k] + 12, f"{tot[k]:.0f} kt\n({100 * tot[k] * 1000 / TRANSPORT_EMISSIONS_T:.3f}%)",
                ha="center", va="bottom", fontsize=8.5)
    ax.set_xticks(x)
    ax.set_xticklabels(SCENARIOS)
    ax.set_ylabel("CO$_2$ reduction (kt per year)")
    ax.set_ylim(0, 640)
    ax.legend(frameon=False, loc="upper left", fontsize=8.5)
    ax.yaxis.grid(True, color="#e6e6e6", lw=0.6)
    ax.set_axisbelow(True)
    fig.savefig(FIGURES / "Figure_4_5_h4_national.png")
    plt.close(fig)


def main():
    plt = plot_style()
    t44, savings, vehicles = table_4_4()
    model, rates, districts = national_model(savings, vehicles)
    t45, _ = table_4_5(model)
    figure_4_4(plt)
    figure_4_5(model, plt)
    return "\n\n".join([
        "## Section 4.5: H4",
        "### Table 4.4",
        t44,
        "### Representativeness of the validation runs",
        representativeness(),
        "### Table 3.4 (inputs of the national scaling model)",
        table_3_4(model, rates, savings, vehicles),
        "### Table 4.5",
        t45,
        f"Share relative to {TRANSPORT_EMISSIONS_LABEL}.",
        "### Sensitivities and checks",
        sensitivities(model, savings, vehicles, districts),
        "Figures: `Figure_4_4_co2_by_class.png`, `Figure_4_5_h4_national.png`.",
    ])


if __name__ == "__main__":
    print_utf8(main())
